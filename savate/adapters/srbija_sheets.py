"""The Serbian federation's own paperwork, in the three shapes it publishes.

savate-srbija.com and savatevojvodina.com put their results out as whatever
their secretary had open at the time: a Word file, a PDF of a letter, once a
whole competition workbook. Three of those shapes carry results and none of
them is a table this archive already reads, so all three are read here. What
they have in common is the federation, not the layout - the same office writes
all of them, in the same two alphabets, with the same vocabulary for a medal and
a weight class, which is what makes one vocabulary enough for all three.

1. THE RANKED WEIGHT-CLASS LIST  (the national championships and cups)

       MLAĐI JUNIORI:
       -60kg:
       1. Petar Jovanov-„Ruma“ Ruma
       2. Miloje Jakovljević-„Sirmijum“ Sr.Mitrovica

   This is the layout `savate_weight_list` already reads out of the federation's
   PDFs, and these documents are the same series in a different container: six
   of them are Word .doc files, which that adapter cannot open because it reads
   PDF word geometry. Here the text is taken from the container - `textutil` for
   .doc, `pdftotext -layout` for PDF - and the same reading is done on lines.

   Nothing about that reading is copied from the other adapter; it is
   re-implemented, because the two must be free to be corrected separately.
   Where the source is an ordinary single-column PDF the two agree row for row,
   which is how this one was checked.

2. THE PROSE TEAM REPORT  (what the national team won abroad)

       Srebrene medalje osvojili su juniori:
       Aleksa Guša do 70kg član KBS "Novi Sad"
       ...
       Zlatnu medalju u kategoriji do 60kg osvojila je izvanredna Kristina
       Džolić, član Savate kluba "Vojvodina" iz Novog Sada

   A letter, not a table: a medal heading, then a sentence per medallist. These
   are the only record the archive has of several editions, but prose is where
   an adapter invents things, so the reading is deliberately narrow and refuses
   far more than it takes. What it refuses is set out under WHAT IS REFUSED.

   A team report is a PARTIAL field by nature - it names the Serbian medallists
   and nobody else - and every read says so in its report. It is not the result
   of the championship; it is what one delegation brought home.

3. THE COMPETITION WORKBOOK  (Loverval 2021, the European assaut championship)

   One Excel file holding an entire championship: a medal table, a sheet per
   weight class with the poule cross-table, the barème, the warnings and the
   final, and three ring running orders that name the red and blue corner of
   every bout. It is the only source in this group that publishes bouts at all,
   so it is the only one here that yields anything but placings.

WHAT IS REFUSED, AND WHY

*A name in an oblique case.* Serbian declines names: "sa Dejanom Ivković" is
Dejana Ivković, "za Kristinu Džolić" is Kristina Džolić. Un-declining them is
guesswork, so a name is only read where the sentence puts it in the nominative -
after a medal, after the award verb, after a bullet, a comma or "i", or
immediately before its weight. A name introduced by a preposition is left where
it lies and counted in the report. Several real medallists are lost that way.
That is the right trade: a name stored in the wrong form is a second person in
the register, and the register is what this archive is for.

*A sentence that claims two different medals.* "... osvojile srebrene medalje
kao i Marina Đekić ... koja je osvojila bronzanu medalju" names three women and
two colours, and nothing in the sentence says which belongs to whom. The whole
sentence is dropped and reported.

*A page whose text layer has lost its spaces.* One PDF here extracts as
"2 srebrenemedaljeosvojilisujunioriVukašinMatović do 60kgi" - the words are
fused, so putting the spaces back would be inventing where they went. Documents
above a threshold of fused words are refused whole, rather than read down to the
few names that happen to have survived.

*Anyone the document says did not medal.* "bez medalje", "nije uspeo", "medalje
nisu osvojili", "plasirao se u finale najesen" - each of those clears the medal
heading that was in force, because otherwise the four athletes listed after
"Uprkos pobedama bez medalje su ostali" become bronze medallists. This is the
single most dangerous thing in the family and it is guarded twice: the phrase
clears the standing medal, and a paragraph that names no weight class clears it
too.

*A gender read off a masculine plural.* "Srebrene medalje osvojili su juniori:"
covers Jovana Vukčević as well as Aleksa Guša - Serbian uses the masculine
plural for a mixed group. So a plural masculine heading sets the age class and
never the gender. A feminine form (juniorka, juniorke, kadetkinje) does set it,
in the singular and the plural, because Serbian does not use it for a mixed
group; a masculine singular attached to one athlete ("senior Marko Vidaković")
sets it too.

WHAT IS READ THAT THE DOCUMENT DOES NOT SPELL OUT, each reported on every read

*Cyrillic is transliterated to Serbian Latin.* Four of these documents are in
Cyrillic and the rest are in Latin, and the same fighters appear in both. Serbian
Cyrillic and Latin are one orthography with two alphabets and a letter-for-letter
mapping between them, so Никола Вукчевић and Nikola Vukčević are not two people
and must not become two rows. The mapping is the standard one, including the
digraphs љ/њ/џ, and it is applied to the whole document before anything else.

*A weight heading with no sign is an upper limit.* "27kg:" between "-24kg:" and
"30kg:" is a dropped hyphen. Every savate class is an upper limit except the open
class, which is always written "+".

*One line carrying both a weight class and an age heading.* The 2021 pioneers'
report prints "-70kg: MLAĐI JUNIORI:" on one line, with three boys under it, and
the previous heading was STARIJE PIONIRKE. Read as a weight heading alone it
files three boys as older pioneer girls - which is exactly what happens without
this - so the line is read as both headings, and the reading is reported.

*A semi-final's winner, where the final says who came out of it.* The workbook
prints the semi-final pairings and the two finalists but not the semi-final
results. A fighter who is in the final won their semi-final; that is a deduction
from the next round, which is what `result_source="inferred"` is for. It is
never blended with a result the document states.

WHAT IS NOT CLAIMED

* No bout is derived from a podium. Every placing here is a placing.
* rouge/bleu is a corner, not a result. Only the workbook's ring running orders
  name corners, and only bouts joined to one of those rows get a corner; the
  rest name a winner and leave the corner empty.
* A competitor's country is filled only where the document gives one. For the
  team reports that is the whole team at once - the document opens "Reprezentacija
  Srbije", so its medallists are Serbian and the report says where that came
  from. The national championships name clubs, never countries, and get none.
* The two savatevojvodina.com tables are not read. They are registers of one
  province's results across eight championships at a time, with no weight class,
  and turning one into a tournament would invent an event that never happened.
  They are recognised and reported as such.
"""

import io
import re
import shutil
import subprocess

from savate import display
from savate import normalize as norm
from savate.schema import MEDALS, Bout, Placing, Report, Tournament

NAME = "srbija_sheets"
DESCRIPTION = ("Serbian federation paperwork: Word ranked lists, prose team "
               "reports, and the Loverval competition workbook")

# ---------------------------------------------------------------- alphabets

# Serbian Cyrillic to Serbian Latin. One orthography, two alphabets, and an
# exact letter-for-letter mapping between them - which is why applying it is
# reading the document rather than translating it. The digraphs are the only
# subtlety: Љ is LJ in a word that is shouting and Lj in one that is not.
_CYRILLIC = {
    "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "ђ": "đ", "е": "e",
    "ж": "ž", "з": "z", "и": "i", "ј": "j", "к": "k", "л": "l", "љ": "lj",
    "м": "m", "н": "n", "њ": "nj", "о": "o", "п": "p", "р": "r", "с": "s",
    "т": "t", "ћ": "ć", "у": "u", "ф": "f", "х": "h", "ц": "c", "ч": "č",
    "џ": "dž", "ш": "š",
}
_CYRILLIC_UPPER = {
    "А": "A", "Б": "B", "В": "V", "Г": "G", "Д": "D", "Ђ": "Đ", "Е": "E",
    "Ж": "Ž", "З": "Z", "И": "I", "Ј": "J", "К": "K", "Л": "L", "Љ": "Lj",
    "М": "M", "Н": "N", "Њ": "Nj", "О": "O", "П": "P", "Р": "R", "С": "S",
    "Т": "T", "Ћ": "Ć", "У": "U", "Ф": "F", "Х": "H", "Ц": "C", "Ч": "Č",
    "Џ": "Dž", "Ш": "Š",
}
_HAS_CYRILLIC = re.compile(r"[Ѐ-ӿ]")


def _latin(text):
    """Serbian Cyrillic -> Serbian Latin, letter for letter."""
    out = []
    chars = str(text or "")
    for i, c in enumerate(chars):
        if c in _CYRILLIC:
            out.append(_CYRILLIC[c])
        elif c in _CYRILLIC_UPPER:
            mapped = _CYRILLIC_UPPER[c]
            nxt = chars[i + 1] if i + 1 < len(chars) else ""
            if len(mapped) == 2 and (nxt in _CYRILLIC_UPPER or nxt.isupper()):
                mapped = mapped.upper()
            out.append(mapped)
        else:
            out.append(c)
    return "".join(out)


def _fold(text):
    """The archive's folding, which already knows Đ is not D with an accent."""
    return display.fold(text)


# ------------------------------------------------------------- vocabularies

# Age class and gender, as this federation writes them in both alphabets (the
# Cyrillic is transliterated before any of this is consulted). The value is
# (age class, gender); "" for gender means the word does not state one, which
# is the case for every masculine plural - Serbian uses it for a mixed group.
_VOCAB = {
    "pionir": ("Pioneer", "Men"), "pioniri": ("Pioneer", ""),
    "pionire": ("Pioneer", ""), "pionira": ("Pioneer", ""),
    "pionirka": ("Pioneer", "Women"), "pionirke": ("Pioneer", "Women"),
    "pionirki": ("Pioneer", "Women"),
    "kadet": ("Cadet", "Men"), "kadeti": ("Cadet", ""),
    "kadete": ("Cadet", ""), "kadeta": ("Cadet", ""),
    "kadetkinja": ("Cadet", "Women"), "kadetkinje": ("Cadet", "Women"),
    "kadetski": ("Cadet", ""), "kadetskom": ("Cadet", ""),
    "junior": ("Junior", "Men"), "juniori": ("Junior", ""),
    "juniore": ("Junior", ""), "juniora": ("Junior", ""),
    "juniorima": ("Junior", ""), "juniorskoj": ("Junior", ""),
    "juniorka": ("Junior", "Women"), "juniorke": ("Junior", "Women"),
    "juniorki": ("Junior", "Women"),
    "senior": ("Senior", "Men"), "seniori": ("Senior", ""),
    "seniore": ("Senior", ""), "seniora": ("Senior", ""),
    "seniorima": ("Senior", ""), "seniorskoj": ("Senior", ""),
    "seniorka": ("Senior", "Women"), "seniorke": ("Senior", "Women"),
    "seniorki": ("Senior", "Women"),
    "veteran": ("Veteran", "Men"), "veterani": ("Veteran", ""),
    "zene": ("", "Women"), "zena": ("", "Women"), "devojke": ("", "Women"),
    "devojcice": ("", "Women"), "dame": ("", "Women"),
    "muskarci": ("", "Men"), "muskaraca": ("", "Men"), "momci": ("", "Men"),
    "men": ("", "Men"), "women": ("", "Women"),
}

# "mlađi juniori" is a class of its own and not a younger sort of junior, so
# the modifier travels with the age word rather than replacing it.
_MODIFIER = {"mladi": "Younger", "mlade": "Younger", "mladji": "Younger",
             "mladje": "Younger", "mladjih": "Younger", "mladja": "Younger",
             "stariji": "Older", "starije": "Older", "starija": "Older",
             "starijih": "Older"}

# Words a heading may carry without being about age or gender.
_IGNORABLE = {"i", "u", "za", "s", "sa", "na", "od", "do", "savate", "savateu",
              "aso", "asaut", "assaut", "combat", "kombat", "boks", "boksu",
              "rezultati", "razultati", "results", "konkurenciji",
              "kategoriji", "kategorija", "mesto", "mjesto", "godine",
              "and", "de", "la", "le"}

# A line that is only a discipline. The federation switches from assaut to
# combat inside one report without repeating the age headings.
_DISCIPLINE = re.compile(r"^(pre[\s-]?combat|combat|kombat|assaut|asaut|aso)"
                         r"(\s+(savate|savateu|finale?|final))?$", re.I)
_DISCIPLINES = {"assaut": "Assaut", "asaut": "Assaut", "aso": "Assaut",
                "combat": "Combat", "kombat": "Combat"}

# Everything after one of these is a club standings board, not results. Those
# rows are numbered too and would otherwise read as placings with a club in
# the fighter column.
# Deliberately only the HEADING of such a table. "Najuspešniji klub na KUPu
# Srbije: ..." is a prize, printed above the results rather than below them,
# and stopping there threw away a whole national cup.
_STOP = re.compile(r"plasman\s+klubova|por[ae]dak\s+klubova|"
                   r"r[ae]zul[ao]?tati\s+klubova|plasman\s+zemalja|"
                   r"medal\s+standing", re.I)

# Different sports, which this federation does also run. A heading naming one
# clears the weight class in force, so its competitors cannot fall into the
# savate class above them.
_OTHER_SPORT = re.compile(r"canne|b[aâ]ton|chausson|forme|\bkanu\b|"
                          r"\bstap\b", re.I)

# A line that says this part of the document IS savate again. Only a heading is
# tested against these - a competitor's club is often "Savate klub Vojvodina",
# and a club name must not be able to switch the sport back on.
# "combat" alone is not enough: CANNE DE COMBAT contains it, and reading that
# as a savate heading puts canne players in a savate weight class.
_SAVATE = re.compile(r"savate|assaut|asaut|\baso\b|"
                     r"(?:combat|kombat)\s+savate", re.I)

# A demonstration is not a bout and its competitors did not place. It has no
# result to record, so it is excluded and counted.
_DEMONSTRATION = re.compile(r"demonstrac|\bdemo\b|egzibicij|exhibition|"
                            r"revijaln|prikaz", re.I)

# The bound a savate class carries, in the words this federation uses for it.
_UNDER_WORDS = {"do", "moins", "under", "ispod"}
_OVER_WORDS = {"preko", "plus", "over", "iznad"}

# Org abbreviations and org words. They are capitalised like a name and sit
# right where a name sits, so a run of capitals starting with one of these is
# a club, not a competitor.
_ORG_TOKENS = {"sk", "sbk", "kbs", "subs", "bk", "su", "kfkb", "kbk", "bkc",
               "sbc", "klub", "kluba", "klubu", "klubovi", "klubova",
               "savate", "boks", "boksersk", "budo", "gym", "sportski",
               "clan", "clanovi", "club"}
_ORG_MARKERS = {"clan", "clanovi", "clanova", "klub", "kluba", "klubu",
                "boks", "boksersk", "budo", "savate"}

# A preposition in front of a name means the name is not in the nominative.
# Serbian declines, and un-declining is guessing, so these block the read.
_OBLIQUE = {"sa", "s", "za", "od", "protiv", "nad", "pod", "prema", "ka",
            "kod", "medju", "medu", "iz", "uz", "posle", "pre", "bez",
            "preko", "nakon", "protivnika", "protivnicu"}

_QUOTES = "„“”\"«»‚‘’"

# ------------------------------------------------------------------ medals

_MEDAL_WORD = re.compile(r"medalj|odlicj|medalja", re.I)
_RANK_COLOURS = [(re.compile(r"\bzlatn", re.I), "1"),
                 (re.compile(r"\bsrebr", re.I), "2"),
                 (re.compile(r"\bbronzan|\bbronz", re.I), "3")]

# The statements that say somebody did NOT win a medal. Each one clears the
# medal heading in force: without that, the athletes listed after "Uprkos
# pobedama bez medalje su ostali" inherit the bronze above them.
_VETO = re.compile(
    r"bez\s+medalj|nisu\s+osvoj|nije\s+osvoj|nije\s+uspe|nisu\s+uspe|"
    r"ostal[aio]\s+bez|ostali\s+bez|medalj\w*\s+nisu|se\s+nisu|"
    r"plasira\w*\s+se\s+u\s+final|nije\s+se\s+plasira|nisu\s+se\s+plasira|"
    r"medalja\s+se\s+nisu", re.I)

# ------------------------------------------------------------------ weights

# A weight class as a heading of its own: "-60kg:", "27kg:", "do 27kg:",
# "+ 85kg:", "M-52kg". The unit is required - without it every "1." and every
# year in a club name becomes a weight class.
_WEIGHT_HEAD = re.compile(
    r"^(?P<pre>[^\d]{0,24}?)"
    r"(?P<sign>[-+–—−])?\s*"
    r"(?<!\d)(?P<a>\d{2,3})(?!\d)"
    r"(?:\s*[-–—/]\s*[-+–—−]?\s*(?<!\d)(?P<b>\d{2,3})(?!\d))?"
    r"\s*(?P<unit>kgs?)\s*[:.]?\s*(?P<rest>.*)$", re.I)

# A weight inside a sentence: "do 65kg", "+70kg", "u kategoriji do 52kg".
_WEIGHT_IN = re.compile(
    r"(?:\b(?P<word>do|preko|ispod|iznad)\s+)?(?P<sign>[-+])?\s*"
    r"(?<!\d)(?P<kg>\d{2,3})(?!\d)\s*kg", re.I)

# A place, printed in front of the competitor.
_RANK = re.compile(r"^(?P<rank>[1-9])(?!\d)\s*[.)°º]?\s*"
                   r"(?:mjesto|mesto|place)?\s*[.:]?\s*(?P<body>\S.*)$", re.I)

_PODIUM = set(MEDALS)

# A four-figure year printed after a club: the competitor's year of birth, on
# three rows of one 2019 cup sheet. It is not part of the club's name.
_BIRTH_YEAR = re.compile(r"\s*\b(19|20)\d{2}\b\s*$")

_MONTHS = {
    "januar": 1, "februar": 2, "mart": 3, "april": 4, "maj": 5, "jun": 6,
    "jul": 7, "avgust": 8, "septembar": 9, "oktobar": 10, "novembar": 11,
    "decembar": 12, "januara": 1, "februara": 2, "marta": 3, "aprila": 4,
    "maja": 5, "juna": 6, "jula": 7, "avgusta": 8, "septembra": 9,
    "oktobra": 10, "novembra": 11, "decembra": 12,
}
_DATE_NUM = re.compile(r"(?<!\d)(\d{1,2})\s*\.\s*(\d{1,2})\s*\.\s*(20\d{2}|19\d{2})(?!\d)")
_DATE_WORD = re.compile(r"(?<!\d)(\d{1,2})\s*\.?\s*([A-Za-zČĆŠŽĐčćšžđ]{3,10})\s+"
                        r"(20\d{2}|19\d{2})(?!\d)")

# "15-16. JUN 2012.", "19-23.07.2023", "09.-12. juna 2016", "21-25. 06. 2023":
# one month, two days. The competition ran from the first to the second.
_MONTH_TOKEN = r"\d{1,2}|[A-Za-zČĆŠŽĐčćšžđ]{3,10}"
_DATE_RANGE_ONE = re.compile(
    r"(?<!\d)(\d{1,2})\s*\.?\s*[-–—]\s*(\d{1,2})\s*\.?\s*"
    r"(" + _MONTH_TOKEN + r")\s*\.?\s*(20\d{2}|19\d{2})(?!\d)", re.I)

# "od 30. avgusta do 02. septembra 2012": a range that crosses a month end.
_DATE_RANGE_TWO = re.compile(
    r"\bod\s+(\d{1,2})\s*\.?\s*([A-Za-zČĆŠŽĐčćšžđ]{3,10})\s+do\s+"
    r"(\d{1,2})\s*\.?\s*([A-Za-zČĆŠŽĐčćšžđ]{3,10})\s+(20\d{2}|19\d{2})(?!\d)",
    re.I)

# The team a report is about. These letters are written by one federation about
# its own delegation, so the delegation is stated once, at the top.
_OUR_TEAM = re.compile(r"reprezentacij\w*\s+srbije|savate\s+savez\w*\s+srbije|"
                       r"sportisti\s+savate\s+saveza\s+srbije|"
                       r"savate\s+saveza\s+srbije", re.I)

# An event heading: a shouted line naming a competition. A document with more
# than one is a federation's annual report holding several championships, and
# then the manifest entry has to say which one it means.
_EVENT_WORDS = re.compile(r"prvenstv|\bkup\b|championship|mastership|"
                          r"izvestaj|rezultati|olimpij", re.I)


def _is_shouted(text):
    letters = [c for c in text if c.isalpha()]
    if len(letters) < 8:
        return False
    return sum(1 for c in letters if c.isupper()) / len(letters) > 0.8


def _tokens(text):
    return [t for t in re.split(r"[^0-9a-z]+", _fold(text)) if t]


def _kg(value):
    return str(value).lstrip("0") or "0"


# -------------------------------------------------------------- the container


class Unreadable(Exception):
    """The document cannot be turned into text at all."""


def _doc_text(path):
    """Plain text out of a Word .doc, through whatever this machine has."""
    for argv in (["textutil", "-convert", "txt", "-stdout", str(path)],
                 ["antiword", str(path)]):
        if not shutil.which(argv[0]):
            continue
        done = subprocess.run(argv, capture_output=True)
        if done.returncode == 0 and done.stdout.strip():
            return done.stdout.decode("utf-8", errors="replace")
    raise Unreadable("no .doc converter on this machine (textutil or antiword)")


def _pdf_text(path):
    if not shutil.which("pdftotext"):
        raise Unreadable("pdftotext (poppler) is required to read a PDF source")
    done = subprocess.run(["pdftotext", "-layout", str(path), "-"],
                          capture_output=True)
    if done.returncode != 0:
        raise Unreadable(f"pdftotext failed: {done.stderr.decode()[:120]}")
    return done.stdout.decode("utf-8", errors="replace")


def _container(path):
    """(kind, payload) for a cached file, from its bytes rather than its name.

    The cache names a file after its URL, so the extension is not reliable and
    for a Wayback capture there is none at all.
    """
    head = path.read_bytes()[:8]
    if head[:4] == b"%PDF":
        return "text", _pdf_text(path)
    if head[:4] == b"\xd0\xcf\x11\xe0":                       # OLE2: Word .doc
        return "text", _doc_text(path)
    if head[:2] == b"PK":                                     # zip: xlsx/docx
        return "workbook", path
    return "text", path.read_text(encoding="utf-8", errors="replace")


_CASE_BREAK = re.compile(r"[a-zà-ÿčćšžđ][A-ZÀ-ÞČĆŠŽĐ]")


def _fused_ratio(text):
    """How much of the document has lost the spaces between its words.

    One PDF in this family extracts as "srebrenemedaljeosvojilisujuniori". A
    word with a capital inside it is the signature; ordinary Serbian has almost
    none, so a document over the threshold is refused rather than guessed at.
    """
    words = [w for w in re.findall(r"[^\W\d_]{3,}", text)]
    if len(words) < 20:
        return 0.0
    return sum(1 for w in words if _CASE_BREAK.search(w)) / len(words)


def _iso(day, month, year):
    """One date, or "" where the numbers are not a date at all."""
    if not (1 <= int(day) <= 31 and 1 <= int(month) <= 12):
        return ""
    return norm.date(f"{int(day):02d}.{int(month):02d}.{year}")


def _month_of(word):
    """The month a date expression names, numeric or written out."""
    text = str(word).strip(". ")
    if text.isdigit():
        return int(text) if 1 <= int(text) <= 12 else 0
    return _MONTHS.get(_fold(text), 0)


def _date_spans(text):
    """[(start, end)] for every date expression in a passage.

    Three shapes, because this federation writes all three: a range inside one
    month ("15-16. JUN 2012.", "09.-12. juna 2016"), a range across two months
    ("od 30. avgusta do 02. septembra 2012"), and a single day ("26.01.2019",
    "18. aprila 2010"). A range is read as a range - taking the last day of one
    and calling it the day the competition started is how this reader used to
    say the Paris world championship began on the 25th when the document says
    it ran on the 24th and the 25th.
    """
    spans, taken = [], []

    def free(found):
        return not any(a < found.end() and found.start() < b for a, b in taken)

    for found in _DATE_RANGE_TWO.finditer(text):
        first, second = _month_of(found.group(2)), _month_of(found.group(4))
        start = _iso(found.group(1), first, found.group(5)) if first else ""
        end = _iso(found.group(3), second, found.group(5)) if second else ""
        if start and end:
            spans.append((start, end))
            taken.append((found.start(), found.end()))
    for found in _DATE_RANGE_ONE.finditer(text):
        month = _month_of(found.group(3))
        start = _iso(found.group(1), month, found.group(4)) if month else ""
        end = _iso(found.group(2), month, found.group(4)) if month else ""
        if start and end and free(found):
            spans.append((start, end))
            taken.append((found.start(), found.end()))
    for pattern in (_DATE_NUM, _DATE_WORD):
        for found in pattern.finditer(text):
            month = _month_of(found.group(2))
            day = _iso(found.group(1), month, found.group(3)) if month else ""
            if day and free(found):
                spans.append((day, ""))
                taken.append((found.start(), found.end()))
    return spans


# A date belongs to a competition where the passage that prints it says so:
# it names a competition, or says one was held. Without that test the first
# date anywhere in the file becomes the competition's - which is how a news
# article's 22:43 timestamp, a letter's signature line and the Tournai leg of
# a different championship all ended up stamped on competitions that the same
# documents date themselves, in so many words, somewhere else.
_DATES_A_COMPETITION = re.compile(
    r"prvenstv|\bkup|championship|mastership|takmicenj|\bigr[ae]\b|odrzan|"
    r"turnir|memorijal", re.I)


def _dates_in(blocks, report=None):
    """(start, end) for the competition these blocks are about.

    Every block that attributes a date to a competition is read, and they have
    to agree: a document that dates one competition twice, differently, has
    contradicted itself and gets no date at all, only a problem saying so.
    """
    found = []
    for index, block in enumerate(blocks):
        spans = _date_spans(block)
        if not spans:
            continue
        attributed = bool(_DATES_A_COMPETITION.search(_fold(block)))
        if not attributed and index and len(block) <= 60:
            # A title line naming the competition, with the date and the town
            # on the line under it. The line under it dates the sheet above.
            previous = blocks[index - 1]
            attributed = bool(_is_shouted(previous)
                              and _DATES_A_COMPETITION.search(_fold(previous)))
        if not attributed:
            continue
        found.append((min(s for s, _e in spans),
                      max(e or s for s, e in spans)))
    if not found:
        return "", ""
    starts = {start for start, _end in found}
    if len(starts) > 1:
        if report is not None:
            report.problem(
                "the document states more than one date for this competition ("
                + ", ".join(sorted(f"{s} to {e}" if e != s else s
                                   for s, e in found)[:4])
                + "); nothing is stored rather than choosing between them")
        return "", ""
    start = starts.pop()
    end = max(end for _start, end in found)
    return start, (end if end and end != start else "")


# ------------------------------------------------------------ shared reading


# A masculine plural means different things in the two layouts, and the
# difference is the document's, not this reader's. A ranked list prints
# MLAĐI JUNIORI and MLAĐE JUNIORKE as two sections of one championship, so
# there the masculine plural is the men's section - the document's own
# opposition says so. A letter about a team writes "osvojili su juniori" over
# a group that includes Jovana Vukčević, because Serbian uses the masculine
# plural for a mixed group. So the list reader reads the gender off it and the
# prose reader never does.
_PLURAL_MEN = {"pioniri", "kadeti", "juniori", "seniori", "veterani", "momci"}


def _read_context(text, sexed_plurals=False):
    """(age, gender, discipline, bound, recognised) from a heading's words."""
    age, gender, discipline, bound, modifier = "", "", "", "", ""
    recognised, said_something = True, False
    for token in _tokens(text):
        if token in _DISCIPLINES:
            discipline = _DISCIPLINES[token]
        elif token in _MODIFIER:
            modifier, said_something = _MODIFIER[token], True
        elif token in _VOCAB:
            klass, who = _VOCAB[token]
            if not who and sexed_plurals and token in _PLURAL_MEN:
                who = "Men"
            age = klass or age
            gender = who or gender
            said_something = True
        elif token in _UNDER_WORDS:
            bound = "under"
        elif token in _OVER_WORDS:
            bound = "over"
        elif token in _IGNORABLE:
            continue
        else:
            recognised = False
    if modifier and age:
        age = f"{modifier} {age}"
    return age, gender, discipline, bound, recognised and said_something


def _single_age(text):
    """The age class a passage names, or "" where it names two of them.

    "za mlađe juniore i kadete" is a European championship and a European cup
    in one sentence, and collapsing the two into one age class - which taking
    the last word seen does - files a cadet as a younger junior.
    """
    ages, modifier = set(), ""
    for token in _tokens(text):
        if token in _MODIFIER:
            modifier = _MODIFIER[token]
        elif token in _VOCAB and _VOCAB[token][0]:
            ages.add(_VOCAB[token][0])
    if len(ages) != 1:
        return ""
    return f"{modifier} {ages.pop()}".strip()


def _is_section(text, sexed_plurals=False):
    """(age, gender, label) if this line is an age/gender heading, else None."""
    body = text.strip().strip(":.-–— ")
    if not body or len(body) > 48 or re.search(r"\d", body):
        return None
    if len(_tokens(body)) > 6:
        return None
    age, gender, _discipline, _bound, recognised = \
        _read_context(body, sexed_plurals)
    if not recognised:
        return None
    return age, gender, " ".join(body.split()).title()


def _bound_of(sign, pre_bound, banded):
    if sign == "+":
        return "over", False
    if sign:
        return "under", False
    if pre_bound:
        return pre_bound, False
    if banded:
        return "under", False
    return "under", True                     # unsigned: read as an upper limit


def _label(parts):
    return " ".join(p for p in parts if p)


def _competitor(body):
    """(name, club) from the text after a place, for the Serbian layout.

    Three separators are in use and sometimes two at once: a dash before a
    quoted club, a quoted club with no dash, and a dash with no quote. The cut
    is taken at the LAST dash before the FIRST quote, because rows routinely
    lose their opening quote in extraction and splitting at the quote alone
    would hand back "Nemanja Grbić-Ruma" as a name - and because a hyphenated
    surname must survive, which splitting at the first dash does not do.
    """
    body = body.replace(" ", " ").strip()
    quote = next((i for i, c in enumerate(body) if c in _QUOTES), -1)
    if quote >= 0:
        head, tail = body[:quote], body[quote:]
    else:
        head, tail = body, ""
    dashes = list(re.finditer(r"\s*[-–—]+\s*", head))
    if dashes:
        cut = dashes[-1]
        name, club = head[:cut.start()], head[cut.end():] + tail
    else:
        name, club = head, tail
    club = "".join(" " if c in _QUOTES else c for c in club)
    club = " ".join(club.strip(" -–—,").split())
    club, cut_year = _BIRTH_YEAR.subn("", club)
    return " ".join(name.strip(" -–—,").split()), club.strip(), bool(cut_year)


# -------------------------------------------------- 1. ranked weight lists


def _read_list(lines, slug, meta, report):
    """Placings from a ranked weight-class list."""
    placings = []
    section_age = meta.get("age_class", "")
    section_gender, section_label, discipline = "", "", ""
    klass = None
    unsigned = unranked = off_podium = orphans = years = 0
    demonstrations = other_sport_rows = 0
    savate = True
    seen = set()

    def place(rank, body):
        nonlocal off_podium, orphans, years
        if rank not in _PODIUM:
            off_podium += 1
            return
        if klass is None:
            orphans += 1
            return
        name, club, had_year = _competitor(body)
        if had_year:
            years += 1
        if len(name) < 3 or not re.search(r"[^\W\d_]{2}", name):
            report.problem(f"{klass['category']}: place {rank} has no readable "
                           f"name in {body[:60]!r} - not stored")
            return
        if re.search(r"\d", name):
            report.problem(f"{klass['category']}: place {rank} reads as "
                           f"{name!r}, which carries digits - not stored")
            return
        key = (klass["category"], rank, _fold(name))
        if key in seen:
            return
        seen.add(key)
        placings.append(Placing(
            tournament=slug, placing_id=f"{slug}-{len(placings) + 1:04d}",
            category=klass["category"], gender=klass["gender"],
            age_class=klass["age_class"], weight_kg=klass["weight_kg"],
            weight_bound=klass["weight_bound"], rank=rank, medal=MEDALS[rank],
            fighter=name, club=club, country="", result_source="reported"))

    def set_class(found, age, gender):
        nonlocal unsigned
        kilos = found.group("b") or found.group("a")
        _a, _g, pre_sport, pre_bound, _ok = _read_context(found.group("pre"),
                                                          True)
        bound, guessed = _bound_of(found.group("sign") or "", pre_bound,
                                   bool(found.group("b")))
        if guessed:
            unsigned += 1
        weight = f"{'+' if bound == 'over' else '-'}{_kg(kilos)} kg"
        head = age or (section_label if not gender else "")
        return {"category": _label((head, gender, pre_sport or discipline,
                                    weight)),
                "gender": gender, "age_class": age,
                "weight_kg": _kg(kilos), "weight_bound": bound}

    for raw in lines:
        line = " ".join(raw.replace(" ", " ").split())
        if not line:
            continue
        if _STOP.search(line):
            report.notes["stopped_at"] = line[:70]
            break
        ranked = bool(_RANK.match(line))
        if not ranked:
            if _OTHER_SPORT.search(line) and not _SAVATE.search(line):
                # Canne de combat, chausson, bâton and savate forme are
                # different sports. Clearing the weight class is not enough -
                # the next weight heading would open a new one under the age
                # heading still in force and file canne players as savate - so
                # the whole section is skipped until a heading says savate.
                savate, klass = False, None
                report.notes.setdefault("other_sport", [])
                if line[:50] not in report.notes["other_sport"]:
                    report.notes["other_sport"].append(line[:50])
                continue
            if _DEMONSTRATION.search(line):
                savate, klass = True, None
                demonstrations += 1
                continue
            if not savate and _SAVATE.search(line):
                savate, klass = True, None
                continue
        if not savate:
            if ranked:
                other_sport_rows += 1
            continue
        if _DISCIPLINE.match(line.strip(":. ")):
            discipline = _DISCIPLINES.get(_fold(line.split()[0]).strip(":."),
                                          "")
            klass = None
            continue

        found = _WEIGHT_HEAD.match(line)
        if found:
            rest = (found.group("rest") or "").strip()
            inner_rank = _RANK.match(rest) if rest else None
            heading = (_is_section(rest, True) if rest and not inner_rank
                       else None)
            if not rest or inner_rank or heading:
                if heading:
                    # "-70kg: MLAĐI JUNIORI:" - the weight class and the next
                    # age heading printed on one line. Both are read, and the
                    # competitors below belong to the age heading.
                    section_age = heading[0] or meta.get("age_class", "")
                    section_gender, section_label = heading[1], heading[2]
                    report.problem(
                        f"one line prints both a weight class and an age "
                        f"heading ({line[:48]!r}); it is read as both, and the "
                        f"competitors under it are filed under the age heading")
                pre_age, pre_gender, _s, _b, _ok = \
                    _read_context(found.group("pre"), True)
                klass = set_class(found, pre_age or section_age,
                                  pre_gender or section_gender)
                if inner_rank:
                    place(inner_rank.group("rank"), inner_rank.group("body"))
                continue

        found = _RANK.match(line)
        if found:
            place(found.group("rank"), found.group("body"))
            continue

        heading = _is_section(line, True)
        if heading:
            section_age = heading[0] or meta.get("age_class", "")
            section_gender, section_label = heading[1], heading[2]
            klass = None
            continue

        if klass is not None and _looks_like_entrant(line):
            unranked += 1

    report.notes["placings"] = len(placings)
    report.notes["categories"] = len({p.category for p in placings})
    if unranked:
        report.problem(f"{unranked} competitor(s) printed under a weight class "
                       f"with no place - listed, not ranked, so not stored")
    if off_podium:
        report.problem(f"{off_podium} placing(s) below third were dropped: this "
                       f"schema stores a podium and fourth is not one")
    if orphans:
        report.problem(f"{orphans} ranked line(s) appeared with no weight class "
                       f"in force and were dropped rather than filed under the "
                       f"previous one")
    if unsigned:
        report.problem(f"{unsigned} weight heading(s) printed no sign; read as "
                       f"an upper limit, which is what every savate class is "
                       f"except the open class, and that is always written '+'")
    if years:
        report.notes["birth_years_dropped"] = years
    if report.notes.get("other_sport"):
        report.problem("heading(s) naming another sport were skipped: " +
                       "; ".join(report.notes["other_sport"][:4]) +
                       " - canne de combat, chausson and savate forme are "
                       "different sports and are not filed as savate")
    if other_sport_rows:
        report.notes["other_sport_rows"] = other_sport_rows
        report.problem(f"{other_sport_rows} ranked line(s) under another "
                       f"sport's heading were dropped; the section runs until "
                       f"a heading names savate again")
    if demonstrations:
        report.notes["demonstrations"] = demonstrations
        report.problem(f"{demonstrations} demonstration bout(s) were excluded: "
                       f"an exhibition has no result and is not a bout")
    _flag_repeated_places(placings, report)
    without_age = {p.category for p in placings if not p.age_class}
    if without_age:
        report.problem(f"{len(without_age)} category label(s) carry no age "
                       f"class because the document states none: " +
                       ", ".join(sorted(without_age)[:4]))
    if not placings:
        report.problem("no ranked weight-class lists found - this is probably "
                       "not a document this reader can read")
    return placings


def _flag_repeated_places(placings, report):
    """Savate awards two bronzes; anything else twice is the document slipping."""
    twice = sorted({p.category for p in placings
                    if p.rank != "3" and sum(
                        1 for q in placings
                        if q.category == p.category and q.rank == p.rank) > 1})
    if twice:
        report.problem(
            "the document awards the same place twice in " +
            ", ".join(twice[:4]) + (" and others" if len(twice) > 4 else "") +
            " - two weight classes printed under one heading, or a place "
            "mistyped; kept as printed rather than reassigned")
    thin = sorted({p.category for p in placings
                   if p.rank == "3" and not any(
                       q.category == p.category and q.rank == "1"
                       for q in placings)})
    if thin:
        report.problem("class(es) " + ", ".join(thin[:4]) + " print a third "
                       "place but no first - kept as printed")


def _looks_like_entrant(text):
    if len(text) > 90 or not re.search(r"[^\W\d_]", text):
        return False
    return bool(re.search(r"[„“”\"]", text) or len(re.split(r"\s{2,}", text)) > 1)


# ----------------------------------------------------- 2. prose team reports


def _blocks(text):
    """The document as blocks: a paragraph, a bullet, or a heading.

    A PDF wraps a paragraph at the right margin and a Word file does not, so
    lines are joined only where one really is a wrap: long, and ending in a
    letter rather than in punctuation. A line ending in a comma or in "i" is
    the federation's list layout, one medallist per line, and joining those
    would put four medallists and four weights into one sentence.
    """
    out, current = [], []
    for raw in text.splitlines():
        line = " ".join(raw.replace(" ", " ").split())
        if not line:
            if current:
                out.append(" ".join(current))
                current = []
            continue
        starts_block = bool(re.match(r"^[-–—•*]\s*\S", line))
        if current and not starts_block:
            previous = current[-1]
            tail = previous.split()[-1] if previous.split() else ""
            # A wrapped line breaks mid-phrase and its continuation begins in
            # lower case. A list line ending in "i" ("... u finalu i") is the
            # federation's "and", and the next medallist's name follows it in
            # capitals - joining those two puts two medallists and one weight
            # into one sentence, and the weight then lands on the wrong one.
            continues = not (len(tail) == 1 and line[:1].isupper())
            if len(previous) >= 78 and previous[-1].isalpha() and continues:
                current.append(line)                       # a wrapped paragraph
                continue
        if current:
            out.append(" ".join(current))
        current = [line]
    if current:
        out.append(" ".join(current))
    return [b.strip() for b in out if b.strip()]


_SENTENCE = re.compile(r"(?<=[.!?])\s+(?=[A-ZÀ-ÞČĆŠŽĐ])")


def _sentences(block):
    return [s.strip() for s in _SENTENCE.split(block) if s.strip()]


def _ranks_in(text):
    """The medal colours a passage states, in the order it states them."""
    hits = []
    for pattern, rank in _RANK_COLOURS:
        found = pattern.search(text)
        if found:
            hits.append((found.start(), rank))
    return [rank for _start, rank in sorted(hits)]


def _weights_in(text):
    return list(_WEIGHT_IN.finditer(text))


def _weight_of(found, report_unsigned):
    word = (found.group("word") or "").lower()
    sign = found.group("sign") or ""
    if sign == "+" or word in _OVER_WORDS:
        bound = "over"
    elif sign == "-" or word in _UNDER_WORDS:
        bound = "under"
    else:
        bound = "under"
        report_unsigned.append(found.group(0))
    return _kg(found.group("kg")), bound


_WORD = re.compile(r"\S+")


def _name_run(sentence, before=None):
    """The competitor a passage names, or (None, reason).

    A run of capitalised words is a name only where the sentence puts it in the
    nominative. Everything about this function is about refusing the rest: a
    club, a town, an org abbreviation, and above all a name governed by a
    preposition, which in Serbian is a declined form this archive must not
    un-decline.
    """
    words = [(m.start(), m.end(), m.group(0)) for m in _WORD.finditer(sentence)]
    if before is not None:
        words = [w for w in words if w[1] <= before]
    runs, current = [], []
    for start, end, word in words:
        core = word.strip(",;:.()-–—•*" + _QUOTES)
        quoted = bool(word) and word[0] in _QUOTES + "(«"
        capital = bool(core) and core[0].isupper() and not re.search(r"\d", core)
        if capital and not quoted and _fold(core) not in _ORG_TOKENS:
            current.append((start, end, core))
            if word[-1] in ",;:." + _QUOTES:
                runs.append(current)
                current = []
        else:
            if current:
                runs.append(current)
            current = []
    if current:
        runs.append(current)

    candidates = [r for r in runs if len(r) >= 2]
    if before is not None:
        # The name a weight belongs to is the one printed immediately in front
        # of it, not the first name in a sentence that lists four of them.
        candidates = [r for r in candidates if 0 <= before - r[-1][1] <= 4]
        candidates = candidates[-1:]
    else:
        candidates = candidates[:1]
    if not candidates:
        return None, "no capitalised pair reads as a name"
    run = candidates[0]

    lead = sentence[:run[0][0]]
    preceding = _WORD.findall(lead)[-3:]
    blocker = ""
    if preceding:
        last = preceding[-1].strip(",;:." + _QUOTES)
        folded = _fold(last)
        if folded in _OBLIQUE:
            blocker = last
        elif folded in _ORG_MARKERS or (last.isupper() and len(last) <= 5
                                        and folded in _ORG_TOKENS):
            blocker = last
        elif preceding[-1] and preceding[-1][-1] in _QUOTES:
            blocker = preceding[-1]
    if blocker:
        return None, f"the name after {blocker!r} is not in the nominative"

    tokens = [t[2] for t in run]
    age, gender = "", ""
    while tokens:
        folded = _fold(tokens[0])
        if folded in _VOCAB:
            klass, who = _VOCAB[folded]
            age, gender = klass or age, who or gender
            tokens.pop(0)
        elif folded in _MODIFIER and len(tokens) > 2:
            age = f"{_MODIFIER[folded]} {age}".strip()
            tokens.pop(0)
        else:
            break
    if len(tokens) < 2:
        return None, "only one word is left once the age class is taken off"

    for word in reversed(preceding):
        folded = _fold(word.strip(",;:."))
        if folded in _VOCAB:
            klass, who = _VOCAB[folded]
            age, gender = klass or age, who or gender
        elif folded in _MODIFIER:
            age = f"{_MODIFIER[folded]} {age}".strip()
    return {"name": " ".join(tokens), "age": age.strip(), "gender": gender,
            "start": run[0][0], "end": run[-1][1]}, ""


def _club_in(text):
    """The club a passage names: its quoted name, with its abbreviation."""
    found = re.search(r"[„“”\"«]\s*([^„“”\"«»]{2,40})\s*[“”\"»]", text)
    if not found:
        return ""
    club = " ".join(found.group(1).split())
    lead = _WORD.findall(text[:found.start()])
    if lead:
        last = lead[-1].strip(",;:.")
        if last.isupper() and 2 <= len(last) <= 5 and last.isalpha():
            return f"{last} {club}"
    return club


# A delegation report that names TWO competitions in the sentence that counts
# its medals. "osvojili su 1 zlatnu, 4 srebrene i 2 bronzane medalje na
# Evropskom prvenstvu za mlađe juniore u asaut savateu i Evropskom KUP-u za
# kadete" is a championship and a cup, and which medal was won at which the
# document never says. The shouted-heading guard cannot see this one, because
# the second competition is named inside a sentence.
_TWO_COMPETITIONS = re.compile(
    r"\bna\s+[^.]{0,80}?(?:prvenstv\w*|\bkup\w*)[^.]{0,80}?\s+i\s+"
    r"[^.]{0,90}?(?:prvenstv\w*|\bkup\w*)", re.I)


def _section(blocks, events, wanted_fold):
    """The blocks that belong to the competition the entry names."""
    if not events or not wanted_fold:
        return list(blocks)
    out, inside = [], False
    for index, block in enumerate(blocks):
        heading = next((e for e in events if e[0] == index), None)
        if heading:
            inside = wanted_fold in _fold(heading[1])
        if inside:
            out.append(block)
    return out


def _two_competitions(blocks):
    """The sentence that credits one set of medals to two competitions."""
    for block in blocks:
        for sentence in _sentences(block):
            if not (_MEDAL_WORD.search(sentence) and _ranks_in(sentence)):
                continue
            found = _TWO_COMPETITIONS.search(sentence)
            if found:
                return " ".join(found.group(0).split())
    return ""


def _flag_unsexed_classes(placings, report):
    """One label, two medallists, and nothing in the document to sex them.

    A letter about a delegation prints a weight and a medal and, most of the
    time, no gender - and the men's -56kg and the women's -56kg are then the
    same string. Two rows under it are either one class's two medallists or
    two classes' one each, and the document does not say which. Stored as
    printed, with the ambiguity said out loud, because inventing the gender
    that would separate them is the one thing this reader may not do.
    """
    groups = {}
    for placing in placings:
        groups.setdefault(placing.category, []).append(placing)
    merged = sorted(label for label, rows in groups.items()
                    if len(rows) > 1 and any(not r.gender for r in rows))
    for label in merged[:4]:
        who = ", ".join(f"{r.fighter} ({r.medal})" for r in groups[label])
        report.problem(
            f"{label or 'a class with no weight'} carries more than one "
            f"medallist and the document states no gender for them, so the "
            f"men's class and the women's class of that weight are one label "
            f"here: {who}. The rows are stored as printed; the podium they "
            f"appear to form may be two competitions' podiums")
    if len(merged) > 4:
        report.notes["merged_labels"] = merged


def _name_variants(text):
    """Every capitalised two-word run the document prints, folded to a key."""
    runs = {}
    for found in re.finditer(r"([A-ZÀ-ÞČĆŠŽĐ][^\W\d_]{1,20})\s+"
                             r"([A-ZÀ-ÞČĆŠŽĐ][^\W\d_]{1,20})", text):
        runs.setdefault(" ".join(found.groups()), 0)
        runs[" ".join(found.groups())] += 1
    return runs


def _one_letter_apart(first, second):
    if abs(len(first) - len(second)) > 1 or first == second:
        return False
    if len(first) > len(second):
        first, second = second, first
    for cut in range(len(second)):
        if first == second[:cut] + second[cut + 1:]:
            return True
    return len(first) == len(second) and sum(
        1 for a, b in zip(first, second) if a != b) == 1


def _flag_name_variants(placings, text, report):
    """A document that spells one competitor's name two ways.

    "Kristina Džolić" in two sentences and "Kristina Đolić" in the third is one
    champion and, stored, two people. Which spelling is right is not in the
    file, so the row keeps the one the sentence that states the medal prints -
    and the reader is told the file disagrees with itself.
    """
    printed = _name_variants(text)
    for placing in placings:
        for other in printed:
            if other == placing.fighter:
                continue
            mine, theirs = placing.fighter.split(), other.split()
            if len(mine) != 2 or len(theirs) != 2 or mine[0] != theirs[0]:
                continue
            if _one_letter_apart(_fold(mine[1]), _fold(theirs[1])):
                report.problem(
                    f"the document spells this competitor's name two ways - "
                    f"{placing.fighter!r} and {other!r} - and does not say "
                    f"which is right; the row keeps the spelling of the "
                    f"sentence that states the medal")


def _read_prose(text, slug, meta, report, options):
    """(placings, (start, end)) from a letter about what the team won."""
    blocks = _blocks(text)
    events = _events(blocks)
    wanted = str(options.get("event") or meta.get("event") or "").strip()
    if events:
        report.notes["events_in_document"] = [e[1] for e in events]
    if len(events) > 1 and not wanted:
        report.problem(
            "this document holds more than one competition and the entry does "
            "not say which: " + "; ".join(e[1] for e in events) +
            ". Add \"event\": \"<one of those>\" to the manifest entry")
        return [], ("", "")
    wanted_fold = _fold(wanted)
    if wanted and not any(wanted_fold in _fold(e[1]) for e in events):
        report.problem(f"the document holds no competition called {wanted!r}; "
                       f"it holds: " + "; ".join(e[1] for e in events))
        return [], ("", "")

    # The date is read from the section that belongs to the competition being
    # read, never from the document at large: the annual report holds three
    # championships, and the first date in the file belongs to one of them.
    mine = _section(blocks, events, wanted_fold)
    dates = _dates_in(mine, report)

    split = _two_competitions(mine)
    if split:
        report.problem(
            f"this document says its medals were won at two competitions "
            f"({split[:90]!r}), and which medal was won at which it does not "
            f"say. Reading them as one event would merge a championship and a "
            f"cup into a competition that never happened, so nothing is read "
            f"from it")
        return [], dates

    ours = "Serbia" if _OUR_TEAM.search(text) else ""
    if ours:
        report.problem(
            "every medallist here is a member of the delegation the document "
            "reports on, which it names in its first sentence; their country "
            "is recorded as Serbia on that statement and on nothing else")

    placings, unsigned = [], []
    standing_rank, standing_age, in_run = "", meta.get("age_class", ""), False
    ambiguous = blocked = weightless = 0
    inside = not events or not wanted
    title_age = _single_age(blocks[0]) if blocks and not events else ""
    standing_age = standing_age or title_age
    for index, block in enumerate(blocks):
        heading = next((e for e in events if e[0] == index), None)
        if heading:
            inside = (not wanted) or (wanted_fold in _fold(heading[1]))
            standing_rank, in_run = "", False
            if inside:
                # "EVROPSKO PRVENSTVO ZA KADETE U SAVATEU U BELGIJI" is a
                # heading that governs its whole section: the medals under it
                # are cadet medals because the heading says so. A heading
                # naming two age classes names none of them.
                standing_age = _single_age(heading[1]) or standing_age
            continue
        if not inside:
            continue
        if _VETO.search(block):
            if standing_rank:
                report.problem(
                    f"the medal heading in force was cleared by {block[:60]!r}, "
                    f"which says these competitors did not medal")
            standing_rank, in_run = "", False
            continue

        ranks = _ranks_in(block)
        weights = _weights_in(block)
        named, _why = _name_run(block)
        # A medal HEADING names a colour and nothing else: "Srebrene medalje
        # osvojili su juniori:". An opening paragraph names colours too, but it
        # COUNTS them - "osvojila je 1 zlatnu, 5 srebrenih i 1 bronzanu" - and
        # it runs on for a paragraph. Reading one of those as a heading hands
        # its medal to whoever the next paragraph happens to mention, which is
        # how a team list became two gold medals.
        heading_shaped = (len(block) <= 90 and not re.search(r"\d", block))
        if (ranks and _MEDAL_WORD.search(block) and not weights and not named
                and heading_shaped):
            standing_rank = ranks[0] if len(ranks) == 1 else ""
            standing_age = _single_age(block) or standing_age
            in_run = False
            continue
        if not ranks and not weights:
            heading = _is_section(block) if len(block) <= 48 else None
            if heading and heading[0]:
                # "juniori:" over one group of medals and "seniori:" over the
                # next is how the annual report splits its age classes.
                standing_age = heading[0]
                continue
            if not named:
                standing_rank, in_run = "", False
                continue

        produced, inline = False, False
        block_rank = standing_rank
        for sentence in _sentences(block):
            sentence_ranks = _ranks_in(sentence)
            if len(sentence_ranks) > 1:
                ambiguous += 1
                report.problem(
                    f"a sentence claims more than one medal colour and names "
                    f"more than one competitor, and nothing in it says which "
                    f"is whose - dropped: {sentence[:80]!r}")
                continue
            # A medal heading governs the paragraphs that follow it; a medal
            # named inside a sentence governs that sentence and stops there.
            # The difference is what keeps "Reprezentaciju Srbije predstavili
            # su takmičari Kristina Džolić do 60kg i Miroslav Odžić do 75kg" -
            # a team list, in a report whose first line mentions a gold - from
            # handing a gold medal to both of them.
            rank = sentence_ranks[0] if sentence_ranks else block_rank
            if sentence_ranks:
                block_rank, inline = sentence_ranks[0], True
            if not rank:
                continue
            found = _weights_in(sentence)
            mentions = []
            if len(found) == 1:
                named, why = _name_run(sentence)
                if named:
                    mentions.append((named, found[0]))
                elif why:
                    blocked += 1
                    report.problem(f"a medal is stated but {why}: "
                                   f"{sentence[:80]!r}")
            elif len(found) > 1:
                for weight in found:
                    named, why = _name_run(sentence, before=weight.start())
                    if named:
                        mentions.append((named, weight))
                    elif why:
                        blocked += 1
            elif (in_run and len(_sentences(block)) == 1
                  and re.search(r"[„“”\"]", sentence)
                  and any(_fold(w) in _ORG_MARKERS or (w.isupper() and 2 <= len(w) <= 5
                          and _fold(w) in _ORG_TOKENS)
                          for w in _WORD.findall(sentence))):
                # A medallist in a list whose weight the document never states.
                # Only inside a run of medallists and only where a club is
                # printed beside the name, because otherwise the sponsors and
                # the coaching staff in the next paragraph read as medallists.
                named, why = _name_run(sentence)
                if named:
                    mentions.append((named, None))
                    weightless += 1
                    report.problem(
                        f"{named['name']} is recorded with no weight class: "
                        f"the document states the medal and not the class")
            # "u konkurenciji kadeta do 52kg" puts the age class after the
            # weight rather than in front of the name. A sentence that names
            # exactly one age class names it for everybody it mentions.
            sentence_age = _single_age(sentence)
            for index, (named, weight) in enumerate(mentions):
                kilos, bound = ("", "")
                if weight is not None:
                    kilos, bound = _weight_of(weight, unsigned)
                # A club is looked for between this competitor and the next
                # one, never past them: in "Ivana Popadić do 52kg iz SK 'Banat'
                # Zrenjanin, Igor Šivoljicki do 70kg član SBK 'Sirmijum'" a
                # competitor printed without a club would otherwise borrow the
                # club of the competitor after them.
                ends = (mentions[index + 1][0]["start"]
                        if index + 1 < len(mentions) else len(sentence))
                tail = sentence[named["end"]:ends]
                age = named["age"] or sentence_age or standing_age
                label = _label((age, named["gender"],
                                f"{'+' if bound == 'over' else '-'}{kilos} kg"
                                if kilos else ""))
                placings.append(Placing(
                    tournament=slug,
                    placing_id=f"{slug}-{len(placings) + 1:04d}",
                    category=label, gender=named["gender"],
                    age_class=age, weight_kg=kilos, weight_bound=bound,
                    rank=rank, medal=MEDALS[rank], fighter=named["name"],
                    club=_club_in(tail), country=ours,
                    result_source="reported"))
                produced = True
        in_run = produced
        if inline or (not produced and not ranks):
            standing_rank = ""

    report.notes["placings"] = len(placings)
    if unsigned:
        report.problem(f"{len(unsigned)} weight(s) were printed with no sign "
                       f"({', '.join(sorted(set(unsigned))[:4])}) and read as "
                       f"an upper limit")
    if blocked:
        report.notes["names_not_in_nominative"] = blocked
    if weightless:
        report.notes["placings_without_a_weight_class"] = weightless
    if ambiguous:
        report.notes["sentences_claiming_two_colours"] = ambiguous
    if placings:
        report.problem(
            "this is one delegation's report of its own medals, so the field "
            "is partial by construction: the competitors of every other nation "
            "are not in the document and are not in these rows. A bronze with "
            "no gold beside it is the normal shape of that, not a fault")
    if not placings:
        report.problem("no medal statement in this document could be read as "
                       "a competitor, a medal and a weight class together")
    _flag_repeated_places(placings, report)
    _flag_unsexed_classes(placings, report)
    _flag_name_variants(placings, text, report)
    return placings, dates


def _events(blocks):
    """[(block index, heading)] for each competition the document announces."""
    events, last = [], None
    for index, block in enumerate(blocks):
        shouted = _is_shouted(block) and _EVENT_WORDS.search(_fold(block))
        if not shouted or _WEIGHT_IN.search(block):
            last = None
            continue
        if last is not None and last == index - 1:
            events[-1] = (events[-1][0], f"{events[-1][1]} {block}")
        else:
            events.append((index, " ".join(block.split())))
        last = index
    return events


# ------------------------------------------------- 3. the Loverval workbook

_CLASS_CODE = re.compile(r"^(?P<sex>[FM])\s*(?P<sign>\+?)\s*(?P<kg>\d{2,3})$")
_FRENCH_MONTHS = {"janvier": 1, "fevrier": 2, "mars": 3, "avril": 4, "mai": 5,
                  "juin": 6, "juillet": 7, "aout": 8, "septembre": 9,
                  "octobre": 10, "novembre": 11, "decembre": 12,
                  "january": 1, "february": 2, "march": 3, "april": 4,
                  "june": 6, "july": 7, "august": 8, "september": 9,
                  "october": 10, "november": 11, "december": 12}
_DAY = re.compile(r"(\d{1,2})\s+([a-z]{3,10})")
_ATHLETES = re.compile(r"(\d+)\s*athl", re.I)
_CHAMPION = re.compile(r"champion\s+d.europe", re.I)
# "Disqualification: - 1 point". The gap must not be allowed to swallow the
# minus sign, or a disqualification scores 1 and every ordinary defeat in the
# workbook reads as one.
_BAREME = re.compile(r"(victoire|egalit|defaite|forfait|disqualification)"
                     r"[^\d-]{0,12}(-?\s?\d)", re.I)


def _cells(sheet):
    """{(row, column): text} for every cell holding something."""
    out = {}
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            text = " ".join(str(cell.value).replace(" ", " ").split())
            if text:
                out[(cell.row, cell.column)] = text
    return out


def _find(cells, predicate):
    return [rc for rc in sorted(cells) if predicate(cells[rc])]


def _country(printed):
    """A nation as the archive spells it, or as the sheet printed it."""
    resolved = display.country(printed)
    return resolved.name if resolved.known else " ".join(printed.split())


def _class_of(code, report):
    found = _CLASS_CODE.match(code.replace(" ", ""))
    if not found:
        return None
    gender = "Women" if found.group("sex") == "F" else "Men"
    bound = "over" if found.group("sign") == "+" else "under"
    kilos = _kg(found.group("kg"))
    return {"gender": gender, "weight_kg": kilos, "weight_bound": bound,
            "category": f"{gender} {'+' if bound == 'over' else '-'}{kilos} kg",
            "code": code}


def _band_class(band, gender):
    """A medal table's weight band -> the class it is.

    "56-60" is the class that ends at 60: a band is written between the class
    below it and its own limit, and the limit is what names the class.
    """
    band = band.replace(" ", "")
    found = re.match(r"^(?P<sign>[-+])?(?P<a>\d{2,3})(?:-(?P<b>\d{2,3}))?$", band)
    if not found:
        return None
    kilos = found.group("b") or found.group("a")
    bound = "over" if found.group("sign") == "+" else "under"
    return {"gender": gender, "weight_kg": _kg(kilos), "weight_bound": bound,
            "category": f"{gender} {'+' if bound == 'over' else '-'}"
                        f"{_kg(kilos)} kg"}


def _entry_register(book, report):
    """{(class code, folded surname): (full name, country)} from 'insc'.

    The class sheets print surnames only; the entry sheet prints the forename
    beside the same surname, in the same class, in the same workbook. Joining
    them is reading the document - and the join is refused wherever one class
    holds two competitors of the same surname.
    """
    if "insc" not in book.sheetnames:
        return {}
    register, clashes = {}, set()
    for row in book["insc"].iter_rows(min_col=1, max_col=5, values_only=True):
        cells = ["" if c is None else " ".join(str(c).split()) for c in row]
        surname, forename, country, code = cells[1], cells[2], cells[3], cells[4]
        if not surname or not code or not _CLASS_CODE.match(code.replace(" ", "")):
            continue
        key = (code, _fold(surname))
        full = " ".join(p for p in (surname, forename) if p)
        if key in register and register[key][0] != full:
            clashes.add(key)
        register[key] = (full, country)
    for key in clashes:
        register.pop(key, None)
        report.problem(f"two competitors named {key[1]!r} are entered in "
                       f"{key[0]} - their forenames are not joined to their bouts")
    return register


def _day_of(header, year):
    """"vendredi 26 novembre / Friday, November the 26th" -> ISO, with the year
    the workbook prints in its own title. Each ring sheet dates its own day,
    and without that two bouts an hour apart on different days both read 09:00."""
    found = _DAY.search(_fold(header))
    if not (found and year):
        return ""
    month = _FRENCH_MONTHS.get(found.group(2))
    if not month:
        return ""
    return norm.date(f"{int(found.group(1)):02d}.{month:02d}.{year}")


def _running_order(book, report, year=""):
    """{(class, phase, the pair): scheduled row} from the ring sheets.

    The ring sheets are the only place in the workbook that states a corner, so
    they are what a corner can come from at all. They cannot be joined on the
    bout number: each ring restarts at 1 on each day and again for the finals,
    so "F48 number 1" is three different bouts. The join is therefore made on
    the class, the round the sheet marks in its own column ("1/2", "F"), and
    the pair of surnames - which is unique inside one class and one round.
    """
    from savate import identity

    order = {}
    for name in book.sheetnames:
        if not re.match(r"^R\d$", name):
            continue
        cells = _cells(book[name])
        ring, day = "", ""
        for (row, column), text in sorted(cells.items()):
            if column == 2 and text.lower().startswith("ring"):
                ring = text
                day = _day_of(cells.get((row, 8), ""), year)
            if column != 1 or not text.isdigit():
                continue
            code = cells.get((row, 3), "")
            red, blue = cells.get((row, 5), ""), cells.get((row, 7), "")
            if not (code and red and blue):
                continue
            marker = cells.get((row, 6), "")
            phase = ("semi" if "1/2" in marker else
                     "final" if marker.strip().upper() == "F" else "poule")
            scheduled = {
                "ring": ring, "time": norm.clock(cells.get((row, 2), "")),
                "date": day,
                "red": red, "blue": blue, "number": text,
                "red_country": cells.get((row, 4), ""),
                "blue_country": cells.get((row, 8), "")}
            for key in _pair_keys(code, phase, red, blue, identity):
                order.setdefault(key, scheduled)
    report.notes["scheduled_bouts"] = len(order)
    return order


def _pair_keys(code, phase, red, blue, identity):
    """The keys one scheduled bout can be found under.

    The running order spells one finalist NANDY and the class sheet spells her
    NANDI. `identity.loose_key` is the archive's own table of the conventions
    that are one name written two ways, so the pair is also filed under that -
    and under nothing looser, because a guess about who fought whom would be
    a guess about who won.
    """
    exact = frozenset((_fold(red), _fold(blue)))
    loose = frozenset((identity.loose_key(red), identity.loose_key(blue)))
    keys = [(code, phase, exact)]
    if loose != exact:
        keys.append((code, phase, loose))
    return keys


def _scheduled_for(order, code, phase, red, blue):
    from savate import identity

    for key in _pair_keys(code, phase, red, blue, identity):
        if key in order:
            return order[key]
    return None


def _year_of(book):
    """The year the workbook prints on its own sheets ("november 2021")."""
    for name in book.sheetnames[:6]:
        for text in _cells(book[name]).values():
            found = re.search(r"\b(19|20)\d{2}\b", text)
            if found and len(text) < 40:
                return found.group(0)
    return ""


def _barometer(book, report):
    """The barème the workbook prints for itself, as {points: meaning}."""
    scale = {}
    if "récap" in book.sheetnames:
        for text in _cells(book["récap"]).values():
            found = _BAREME.search(_fold(text))
            if found:
                word = _fold(found.group(1))
                points = int(found.group(2).replace(" ", ""))
                scale[points] = {"victoire": "win", "egalit": "draw",
                                 "defaite": "loss", "forfait": "forfait",
                                 "disqualification": "disqualification"
                                 }.get(word[:13], word)
    if not scale:
        scale = {3: "win", 1: "loss", 0: "forfait", -1: "disqualification"}
        report.problem("the workbook does not print its barème; the archive's "
                       "own (3 / 1 / 0 / -1) was used to read its scorelines")
    else:
        report.notes["bareme"] = {str(k): v for k, v in sorted(scale.items())}
    return scale


def _points_columns(cells, header_row):
    """(points columns, warning columns) as {column: surname}.

    The header prints every competitor twice, once over the points block and
    once over the warnings block, in the same order - so the split is where the
    first name comes round again.
    """
    named = [(column, cells[(header_row, column)])
             for (row, column) in sorted(cells) if row == header_row
             and column > 3 and cells[(row, column)].upper() != "TIREUR"]
    points, warnings, seen = {}, {}, {}
    half = None
    for index, (column, name) in enumerate(named):
        if name in seen and half is None:
            half = index
        seen[name] = column
    for index, (column, name) in enumerate(named):
        if half is not None and index >= half:
            warnings[column] = name
        else:
            points[column] = name
    return points, warnings


def _score(cells, row, columns, name):
    for column, who in columns.items():
        if who == name and (row, column) in cells:
            value = cells[(row, column)]
            try:
                return int(float(value.replace(",", ".")))
            except ValueError:
                return None
    return None


def _outcome(red_points, blue_points, red_warnings, blue_warnings, scale):
    """(winner corner, decision, status) from a poule scoreline."""
    if red_points is None and blue_points is None:
        return "", "", "unresolved"
    if red_points is None or blue_points is None:
        # One cell filled and one blank. The barème makes 3 a victory outright,
        # so the sheet does say who won - and says nothing at all about how,
        # which is why the decision stays empty rather than being called points.
        scored = red_points if blue_points is None else blue_points
        if scale.get(scored) == "win":
            return ("red" if blue_points is None else "blue"), "", "decided"
        return "", "", "unresolved"
    meaning = {scale.get(red_points, ""), scale.get(blue_points, "")}
    if red_points == blue_points:
        return "", "draw" if "draw" in meaning else "", "unresolved"
    corner = "red" if red_points > blue_points else "blue"
    losing = min(red_points, blue_points)
    decision = {"loss": "points", "forfait": "forfait",
                "disqualification": "disqualification",
                "draw": ""}.get(scale.get(losing, ""), "")
    if not decision:
        loser_warnings = blue_warnings if corner == "red" else red_warnings
        if loser_warnings is not None and loser_warnings >= 3:
            decision = "disqualification"
    return corner, decision, "decided"


def _read_workbook(path, slug, report):
    """Placings and bouts from one competition workbook."""
    try:
        import openpyxl
    except ImportError:                                    # pragma: no cover
        report.problem("openpyxl is required to read this workbook")
        return [], []
    import warnings as _warnings
    with _warnings.catch_warnings():
        _warnings.simplefilter("ignore")
        book = openpyxl.load_workbook(io.BytesIO(path.read_bytes()),
                                      data_only=True)
    report.notes["sheets"] = len(book.sheetnames)
    scale = _barometer(book, report)
    register = _entry_register(book, report)
    year = _year_of(book)
    order = _running_order(book, report, year)

    rows = []
    placings = _medal_table(book, slug, report)
    for name in book.sheetnames:
        if _CLASS_CODE.match(name.replace(" ", "")):
            rows += _class_sheet(book[name], name, slug, register, order,
                                 scale, report)
    if "FIN" in book.sheetnames:
        rows += _finals_sheet(book["FIN"], slug, register, order, report)

    counter = {}
    for bout in rows:
        key = (bout.category, bout.phase)
        counter[key] = counter.get(key, 0) + 1
        bout.bout_id = (f"{slug}-{bout.category.replace(' ', '')}"
                        f"-{bout.phase}-{counter[key]:02d}")
    report.notes["bouts"] = len(rows)
    report.notes["placings"] = len(placings)
    _flag_repeated_places(placings, report)
    _cross_check(placings, rows, report)
    return placings, rows


def _medal_table(book, slug, report):
    """Placings from the CL sheet: four medal columns, two of them bronze."""
    if "CL" not in book.sheetnames:
        report.problem("the workbook has no CL sheet, so it states no podium")
        return []
    cells = _cells(book["CL"])
    columns = [(4, 5, "1"), (7, 8, "2"), (10, 11, "3"), (13, 14, "3")]
    placings, gender = [], ""
    for row in sorted({r for r, _c in cells}):
        band = cells.get((row, 2), "")
        text = " ".join(cells.get((row, c), "") for c in range(1, 15))
        if re.search(r"masculin|\bmen\b", text, re.I):
            gender = "Men"
            continue
        if re.search(r"f[ée]minin|\bwomen\b", text, re.I):
            gender = "Women"
            continue
        klass = _band_class(band, gender) if band else None
        if not klass or not gender:
            continue
        for name_column, country_column, rank in columns:
            printed = cells.get((row, name_column), "")
            if not printed or len(printed) < 3:
                continue
            placings.append(Placing(
                tournament=slug, placing_id=f"{slug}-{len(placings) + 1:04d}",
                category=klass["category"], gender=klass["gender"],
                age_class="", weight_kg=klass["weight_kg"],
                weight_bound=klass["weight_bound"], rank=rank,
                medal=MEDALS[rank], fighter=" ".join(printed.split()),
                country=_country(cells.get((row, country_column), "")),
                club="", result_source="reported"))
        entered = cells.get((row, 1), "")
        if entered.isdigit():
            report.notes.setdefault("entered", {})[klass["category"]] = entered
    return placings


def _finalists(cells, report, code):
    """(the two finalists, the winner) from a sheet's finale block."""
    finale = _find(cells, lambda t: t.strip().lower() == "finale")
    names = []
    if finale:
        row, column = finale[0]
        # The two finalists are the names in the column the "Finale" label
        # heads, which is one to its left in every sheet in this workbook.
        names = [cells[(r, c)] for (r, c) in sorted(cells)
                 if c == column - 1 and r > row
                 and not re.match(r"^n°|^pays$", cells[(r, c)], re.I)
                 and not display.country(cells[(r, c)]).known]
        seen, ordered = set(), []
        for name in names:
            if name not in seen:
                seen.add(name)
                ordered.append(name)
        names = ordered
    champion = _find(cells, lambda t: bool(_CHAMPION.search(t)))
    winner = ""
    if champion:
        row, column = champion[0]
        winner = cells.get((row - 1, column - 1), "")
    if winner and names and winner not in names:
        report.problem(f"{code}: the champion {winner!r} is not one of the two "
                       f"names in the final block ({', '.join(names[:2])}) - "
                       f"the final is stored with no winner")
        winner = ""
    return names[:2], winner


def _class_sheet(sheet, code, slug, register, order, scale, report):
    """Bouts from one weight class's sheet: its poules, semis and final."""
    cells = _cells(sheet)
    klass = _class_of(code, report)
    if not klass:
        return []

    countries = {}
    for (row, column), text in sorted(cells.items()):
        if column == 1 and re.match(r"^tireur\s*\d+$", text, re.I):
            surname = cells.get((row, 2), "")
            for offset in (3, 4, 5):
                printed = cells.get((row, offset), "")
                if printed and not printed.isdigit():
                    countries[surname] = printed
                    break

    def fighter(surname):
        full, country = register.get((code, _fold(surname)), ("", ""))
        printed = countries.get(surname, country)
        return (full or surname, _country(printed) if printed else "")

    bouts, poule_label = [], ""
    headers = {row: _points_columns(cells, row) for (row, column), text
               in cells.items() if column == 2 and text.upper() == "TIREUR"}
    for (row, column), text in sorted(cells.items()):
        if column != 1:
            continue
        label = re.search(r"[Pp]oule\s+([A-Z])\b", text)
        if label:
            poule_label = label.group(1).upper()
        if not re.match(r"^assaut", text, re.I):
            continue
        header = max((r for r in headers if r < row), default=None)
        points, warnings = headers.get(header, ({}, {}))
        red, blue = cells.get((row, 2), ""), cells.get((row, 3), "")
        if not (red and blue):
            continue
        if red == blue:
            report.problem(f"{code}: a bout row pairs {red!r} with themselves "
                           f"- not stored")
            continue
        red_points = _score(cells, row, points, red)
        blue_points = _score(cells, row, points, blue)
        red_warnings = _score(cells, row, warnings, red)
        blue_warnings = _score(cells, row, warnings, blue)
        number = cells.get((row, 4), "")
        scheduled = _scheduled_for(order, code, "poule", red, blue)
        bouts.append(_bout(slug, klass, poule_label, "poule", red, blue,
                           fighter, red_points, blue_points, red_warnings,
                           blue_warnings, scale, scheduled))

    semis = [rc for rc in sorted(cells)
             if rc[1] == 1 and re.match(r"^1/2\s*finale", cells[rc], re.I)]
    pairs = []
    if semis:
        row = semis[0][0]
        while (row, 2) in cells and (row, 4) in cells:
            pairs.append((cells[(row, 2)], cells[(row, 4)]))
            row += 1
    names, winner = _finalists(cells, report, code)
    for red, blue in pairs:
        through = [n for n in (red, blue) if n in names]
        won = through[0] if len(through) == 1 else ""
        bouts.append(_bout(slug, klass, "", "semi", red, blue, fighter,
                           None, None, None, None, scale,
                           _scheduled_for(order, code, "semi", red, blue),
                           winner=won, source="inferred" if won else ""))
    if len(names) == 2:
        bouts.append(_bout(slug, klass, "", "final", names[0], names[1],
                           fighter, None, None, None, None, scale,
                           _scheduled_for(order, code, "final", *names),
                           winner=winner, source="reported" if winner else ""))
    if pairs:
        report.notes["semi_finals_inferred"] = \
            report.notes.get("semi_finals_inferred", 0) + sum(
                1 for red, blue in pairs
                if len([n for n in (red, blue) if n in names]) == 1)
    return bouts


def _finals_sheet(sheet, slug, register, order, report):
    """The finals of the classes that were a final and nothing else."""
    cells = _cells(sheet)
    codes = _find(cells, lambda t: bool(_CLASS_CODE.match(t.replace(" ", ""))))
    bouts = []
    for index, (row, column) in enumerate(codes):
        end = codes[index + 1][0] if index + 1 < len(codes) else 10 ** 6
        block = {rc: t for rc, t in cells.items() if row <= rc[0] < end}
        code = cells[(row, column)]
        klass = _class_of(code, report)
        if not klass:
            continue
        names = [block[rc] for rc in sorted(block)
                 if rc[1] == 2 and re.match(r"^tireur", block.get((rc[0], 1), ""),
                                            re.I)]
        _pair, winner = _finalists(block, report, code)
        if len(names) != 2:
            report.problem(f"{code}: the finals sheet does not name two "
                           f"finalists - no bout stored")
            continue
        if winner and winner not in names:
            report.problem(f"{code}: the champion {winner!r} is not one of "
                           f"{names} - the final is stored with no winner")
            winner = ""

        def fighter(surname, code=code):
            full, country = register.get((code, _fold(surname)), ("", ""))
            return (full or surname, _country(country) if country else "")

        bouts.append(_bout(slug, klass, "", "final", names[0], names[1],
                           fighter, None, None, None, None, {},
                           _scheduled_for(order, code, "final", *names),
                           winner=winner, source="reported" if winner else ""))
    return bouts


def _bout(slug, klass, poule, phase, red, blue, fighter, red_points,
          blue_points, red_warnings, blue_warnings, scale, scheduled,
          winner=None, source=None):
    """One bout row, with the corner filled only where a sheet states one."""
    red_name, red_country = fighter(red)
    blue_name, blue_country = fighter(blue)
    corner, decision, status = "", "", "unresolved"
    result_source = ""
    won, lost = "", ""
    if winner is None:
        corner, decision, status = _outcome(
            red_points, blue_points, red_warnings, blue_warnings, scale)
        if corner:
            won = red_name if corner == "red" else blue_name
            lost = blue_name if corner == "red" else red_name
            result_source = "reported"
    elif winner:
        won = fighter(winner)[0]
        lost = blue_name if won == red_name else red_name
        status, result_source = "decided", source or "reported"
    # The corner is a fact about where a fighter stood, which only the ring
    # sheets state. Where one does, it is used - and if it names the pair the
    # other way round, the row is turned round with it rather than relabelled.
    ring, when, on = "", "", ""
    if scheduled:
        ring, when, on = scheduled["ring"], scheduled["time"], scheduled["date"]
        if scheduled["blue"] == red:
            red_name, blue_name = blue_name, red_name
            red_country, blue_country = blue_country, red_country
            red_points, blue_points = blue_points, red_points
            red_warnings, blue_warnings = blue_warnings, red_warnings
            corner = {"red": "blue", "blue": "red"}.get(corner, corner)
    elif winner is None:
        corner = ""                       # no sheet says which corner was which
    if not scheduled:
        corner = ""
    elif won and not corner:
        corner = "red" if won == red_name else "blue" if won == blue_name else ""
    if corner and won and won != (red_name if corner == "red" else blue_name):
        corner = ""
    return Bout(
        tournament=slug,
        # Numbered per class and round in _read_workbook, where the whole set
        # is known; the workbook's own numbers restart on every ring and every
        # day and so cannot be an identifier.
        bout_id="",
        category=klass["category"], gender=klass["gender"],
        weight_kg=klass["weight_kg"], weight_bound=klass["weight_bound"],
        phase=phase, poule=poule, ring=ring, time=when, date=on,
        red=red_name, red_country=red_country,
        blue=blue_name, blue_country=blue_country,
        red_points="" if red_points is None else str(red_points),
        blue_points="" if blue_points is None else str(blue_points),
        red_warnings="" if red_warnings is None else str(red_warnings),
        blue_warnings="" if blue_warnings is None else str(blue_warnings),
        winner_corner=corner, winner=won, loser=lost,
        decision=decision if winner is None else "",
        status=status, result_source=result_source)


def _cross_check(placings, bouts, report):
    """Does the medal table agree with the finals the sheets print?"""
    golds = {p.category: p.fighter for p in placings if p.rank == "1"}
    for bout in bouts:
        if bout.phase != "final" or not bout.winner:
            continue
        gold = golds.get(bout.category, "")
        if gold and _fold(gold) != _fold(bout.winner):
            report.problem(
                f"{bout.category}: the medal table gives gold to {gold!r} but "
                f"the class sheet's final was won by {bout.winner!r} - both "
                f"are kept as printed")


# ---------------------------------------------------------------- the reader


_REGISTER = re.compile(r"\bpresime\s+i\s+ime\b|\bprezime\s+i\s+ime\b", re.I)


def read(source, slug, meta=None, **options):
    """(Tournament, [Placing|Bout], Report) from one Serbian federation file."""
    from savate import sources

    meta = dict(meta or {})
    report = Report(source=str(source), adapter=NAME)
    tournament = Tournament(
        slug=slug, name=meta.get("name", slug),
        discipline=meta.get("discipline", ""), level=meta.get("level", ""),
        format=meta.get("format", ""), age_class=meta.get("age_class", ""),
        year=meta.get("year", ""), country=meta.get("country", ""),
        city=meta.get("city", ""), start_date=meta.get("start_date", ""),
        end_date=meta.get("end_date", ""), source=str(source), adapter=NAME)

    try:
        path = (sources.fetch_archived(str(source)) if options.get("archived")
                else sources.fetch(source, refresh=options.get("refresh", False)))
    except Exception as e:                              # network, 404, 403, type
        report.problem(f"could not fetch the document: {e}")
        return tournament, [], report

    try:
        kind, payload = _container(path)
    except Unreadable as e:
        report.problem(f"could not turn the document into text: {e}")
        return tournament, [], report
    except Exception as e:                              # pragma: no cover
        report.problem(f"could not read the document: {e}")
        return tournament, [], report

    if kind == "workbook":
        report.notes["mode"] = "workbook"
        try:
            placings, bouts = _read_workbook(payload, slug, report)
        except Exception as e:                          # pragma: no cover
            report.problem(f"could not read the workbook: {e}")
            return tournament, [], report
        # The workbook dates itself: each ring sheet heads its own day, and the
        # bouts carry those dates, so the competition ran from the first to the
        # last of them. Nothing here is filled in over a manifest entry.
        days = sorted({b.date for b in bouts if b.date})
        if days:
            tournament.start_date = tournament.start_date or days[0]
            tournament.end_date = tournament.end_date or days[-1]
            tournament.year = tournament.year or days[0][:4]
        report.read = len(placings) + len(bouts)
        return tournament, placings + bouts, report

    text = payload
    if _HAS_CYRILLIC.search(text):
        text = _latin(text)
        report.notes["alphabet"] = "Cyrillic, transliterated to Serbian Latin"
        report.problem("this document is written in Serbian Cyrillic; its "
                       "names, clubs and headings are transliterated letter "
                       "for letter into Serbian Latin, which is the same "
                       "orthography in the other alphabet - otherwise the same "
                       "fighter is two people in the register")

    fused = _fused_ratio(text)
    if fused > 0.08:
        report.notes["fused_words"] = round(fused, 3)
        report.problem(
            f"{fused:.0%} of this document's words have lost the spaces "
            f"between them in its text layer ('srebrenemedaljeosvojilisujuniori'). "
            f"Where the words went is not in the file, so nothing is read from "
            f"it rather than guessing at the few names that survived")
        return tournament, [], report

    if _REGISTER.search(text) and re.search(r"broj\s+pobeda|broj\s+ucesnika",
                                            _fold(text)):
        report.problem(
            "this is a register of one province's results across several "
            "championships at a time - one row per athlete-result, with no "
            "weight class - and not one competition. Reading it as a tournament "
            "would invent an event, so nothing is read from it")
        return tournament, [], report

    lines = text.splitlines()
    report.read = len(lines)
    mode = options.get("mode") or _shape(lines)
    report.notes["mode"] = mode
    if mode == "list":
        rows = _read_list(lines, slug, meta, report)
        dates = _dates_in(_blocks(text), report)
    else:
        rows, dates = _read_prose(text, slug, meta, report, options)

    start, end = dates
    if start:
        report.notes["date_in_document"] = start if not end else f"{start}/{end}"
    if start and not tournament.start_date:
        tournament.start_date = start
    if end and not tournament.end_date:
        tournament.end_date = end
    if start and not tournament.year:
        tournament.year = start[:4]
    return tournament, rows, report


def _shape(lines):
    """Which of the two text layouts this document is in.

    A ranked list is decided on its own evidence - numbered places under weight
    headings - and never on the file's name or its extension.
    """
    ranks = sum(1 for line in lines if _RANK.match(line.strip()))
    weights = sum(1 for line in lines
                  if _WEIGHT_HEAD.match(" ".join(line.split())))
    return "list" if ranks >= 4 and weights >= 2 else "prose"
